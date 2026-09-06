#!/usr/bin/env python3
"""Reconcile books.csv + gstr2b.json and write a working paper with a 4th
"AI Explanations" sheet: for every HIGH/MEDIUM-severity finding, a cited,
LLM-written explanation grounded in the actual CGST Act/Rules text (or an
explicit "no supporting reference located" if nothing cleared the
similarity floor -- see rag/generation/explain.py).

Requires Qdrant reachable and OPENAI_API_KEY set, and the corpus already
embedded (scripts/embed_corpus.py). Costs one embedding call and, for each
grounded explanation, one chat completion call.
"""

from __future__ import annotations

import argparse
import asyncio
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from gstrecon import ingest, report
from gstrecon.exceptions import DEFAULT_TAX_TOLERANCE, reconcile_full
from gstrecon.matching import DEFAULT_GSTIN_FUZZY_THRESHOLD, DEFAULT_INVOICE_FUZZY_THRESHOLD
from gstrecon.rag.generation.explain import explain_findings
from gstrecon.rag.generation.report_integration import write_explanations_sheet


async def main_async(books: Path, gstr2b: Path, output: Path, tax_tolerance: Decimal) -> None:
    # Same "bad input, not a bug" distinction as cli.py: SchemaMismatchError
    # (missing columns/keys) and normalize.py's plain ValueError (an
    # unparseable date/amount cell) both need a clean message here, not a
    # raw traceback -- a script is not exempt from the standard just
    # because it isn't the primary `gstrecon` CLI entry point.
    try:
        books_docs = ingest.parse_books_csv(books)
        gstr2b_docs = ingest.parse_gstr2b_json(gstr2b)
    except ValueError as exc:
        raise SystemExit(f"Input format error: {exc}") from exc
    reconciliation = reconcile_full(books_docs, gstr2b_docs, tax_tolerance=tax_tolerance)

    metadata = report.RunMetadata(
        books_path=books,
        gstr2b_path=gstr2b,
        tax_tolerance=tax_tolerance,
        invoice_fuzzy_threshold=DEFAULT_INVOICE_FUZZY_THRESHOLD,
        gstin_fuzzy_threshold=DEFAULT_GSTIN_FUZZY_THRESHOLD,
        books_document_count=len(books_docs),
        gstr2b_document_count=len(gstr2b_docs),
    )
    bridge = report.write_working_paper(
        output,
        books_docs=books_docs,
        gstr2b_docs=gstr2b_docs,
        reconciliation=reconciliation,
        metadata=metadata,
    )
    print(f"Wrote {output} (3 sheets). Closing variance: {bridge.closing_variance}")

    print(f"Generating explanations for {len(reconciliation.findings)} findings...")
    explanations = await explain_findings(reconciliation.findings)
    in_scope = sum(1 for e in explanations if e is not None)
    grounded = sum(1 for e in explanations if e is not None and e.grounded)
    print(f"  {in_scope} in scope, {grounded} grounded with a citation")

    workbook = load_workbook(output)
    write_explanations_sheet(workbook.create_sheet(), reconciliation.findings, explanations)
    workbook.save(output)
    print(f"Wrote {output} (4 sheets, incl. AI Explanations).")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("books", type=Path)
    parser.add_argument("gstr2b", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=Path("working_paper.xlsx"))
    parser.add_argument("--tax-tolerance", type=str, default=str(DEFAULT_TAX_TOLERANCE))
    args = parser.parse_args()
    # decimal.InvalidOperation does not subclass ValueError, so argparse's
    # own type= conversion (which only catches ValueError/TypeError) would
    # NOT turn a garbage --tax-tolerance into a clean argparse error --
    # verified this fails with a raw traceback before adding this
    # try/except, the same bug class already found and fixed in cli.py.
    try:
        tax_tolerance = Decimal(args.tax_tolerance)
    except InvalidOperation:
        parser.error(f"--tax-tolerance: invalid decimal value: {args.tax_tolerance!r}")
    asyncio.run(main_async(args.books, args.gstr2b, args.output, tax_tolerance))


if __name__ == "__main__":
    main()
