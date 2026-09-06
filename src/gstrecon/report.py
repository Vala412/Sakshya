"""The working paper: an ITC bridge that ties books to GSTR-2B to the paisa,
an exception register, and a reproducibility sheet, written to a single
Excel workbook.

The ITC bridge is built as a pure arithmetic decomposition of
`Document.signed_tax` -- not by hand-assigning a sign to each reason code --
because a decomposition is correct by construction: every document
contributes exactly once to exactly one bridge line, using the same
`signed_tax` the rest of the engine already trusts, so the closing variance
is provably zero rather than "zero because nobody found a counterexample
yet." `tests/test_report.py::test_bridge_ties_to_nil_property` checks this
against randomized documents specifically to keep that guarantee honest.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from gstrecon import __version__
from gstrecon.exceptions import Exception_, ReconciliationResult
from gstrecon.matching import MatchCategory
from gstrecon.models import Document, ItcAvailability

_ZERO = Decimal("0.00")

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TOTAL_FONT = Font(bold=True)


@dataclass(frozen=True)
class BridgeLine:
    label: str
    reason_codes: str
    document_count: int
    amount: Decimal


@dataclass(frozen=True)
class ItcBridge:
    books_total: Decimal
    gstr2b_total_available: Decimal
    gstr2b_total_blocked: Decimal
    lines: list[BridgeLine]
    closing_variance: Decimal


def _effective_2b_signed_tax(doc: Document) -> Decimal:
    """A 2B document only counts toward "ITC as per 2B" if 2B itself marks
    it available -- an EX-09-flagged document contributes zero here, which
    is exactly what makes the bridge (not just the register) reflect that
    the credit is blocked."""
    return _ZERO if doc.itc_availability is ItcAvailability.NOT_AVAILABLE else doc.signed_tax


def build_itc_bridge(
    books_docs: list[Document],
    gstr2b_docs: list[Document],
    reconciliation: ReconciliationResult,
) -> ItcBridge:
    books_total = sum((d.signed_tax for d in books_docs), _ZERO)
    gstr2b_total_available = sum((_effective_2b_signed_tax(d) for d in gstr2b_docs), _ZERO)
    gstr2b_total_blocked = sum(
        (d.total_tax for d in gstr2b_docs if d.itc_availability is ItcAvailability.NOT_AVAILABLE),
        _ZERO,
    )

    books_only = [
        r.books_doc
        for r in reconciliation.match_results
        if r.category is MatchCategory.IN_BOOKS_NOT_2B and r.books_doc is not None
    ]
    gstr2b_only = [
        r.gstr2b_doc
        for r in reconciliation.match_results
        if r.category is MatchCategory.IN_2B_NOT_BOOKS and r.gstr2b_doc is not None
    ]
    matched_pairs: list[tuple[Document, Document]] = [
        (r.books_doc, r.gstr2b_doc)
        for r in reconciliation.match_results
        if r.books_doc is not None and r.gstr2b_doc is not None
    ]

    lines = [
        BridgeLine(
            label="In books, not in GSTR-2B",
            reason_codes="EX-01 / EX-12",
            document_count=len(books_only),
            amount=sum((-doc.signed_tax for doc in books_only), _ZERO),
        ),
        BridgeLine(
            label="Duplicate booking (books side)",
            reason_codes="EX-10",
            document_count=len(reconciliation.books_duplicate_surplus),
            amount=sum(
                (-doc.signed_tax for doc in reconciliation.books_duplicate_surplus), _ZERO
            ),
        ),
        BridgeLine(
            label="In GSTR-2B, not in books",
            reason_codes="EX-02 / EX-11",
            document_count=len(gstr2b_only),
            amount=sum((_effective_2b_signed_tax(doc) for doc in gstr2b_only), _ZERO),
        ),
        BridgeLine(
            label="Duplicate booking (2B side)",
            reason_codes="EX-10",
            document_count=len(reconciliation.gstr2b_duplicate_surplus),
            amount=sum(
                (_effective_2b_signed_tax(doc) for doc in reconciliation.gstr2b_duplicate_surplus),
                _ZERO,
            ),
        ),
        BridgeLine(
            label="Tax/value variance on matched pairs (see Exception Register)",
            reason_codes="EX-03 / EX-04 / EX-05 / EX-06 / EX-09",
            document_count=len(matched_pairs),
            amount=sum(
                (_effective_2b_signed_tax(g) - b.signed_tax for b, g in matched_pairs), _ZERO
            ),
        ),
    ]

    closing_variance = (
        books_total + sum((line.amount for line in lines), _ZERO) - gstr2b_total_available
    )

    return ItcBridge(
        books_total=books_total,
        gstr2b_total_available=gstr2b_total_available,
        gstr2b_total_blocked=gstr2b_total_blocked,
        lines=lines,
        closing_variance=closing_variance,
    )


@dataclass(frozen=True)
class RunMetadata:
    """Everything needed to reproduce this exact working paper from the same
    two source files: what they were, what settings were used, and when it
    ran. Without this, "audit-defensible" is just an assertion -- a reviewer
    re-running the engine months later needs to confirm they're feeding it
    the same inputs and settings, not guessing.
    """

    books_path: Path
    gstr2b_path: Path
    tax_tolerance: Decimal
    invoice_fuzzy_threshold: float
    gstin_fuzzy_threshold: float
    books_document_count: int
    gstr2b_document_count: int
    generated_at: datetime | None = None

    @staticmethod
    def _sha256(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def to_rows(self) -> list[tuple[str, str]]:
        generated_at = self.generated_at or datetime.now(UTC)
        return [
            ("gstrecon version", __version__),
            ("Generated at (UTC)", generated_at.isoformat(timespec="seconds")),
            ("Books source file", str(self.books_path)),
            ("Books file SHA-256", self._sha256(self.books_path)),
            ("GSTR-2B source file", str(self.gstr2b_path)),
            ("GSTR-2B file SHA-256", self._sha256(self.gstr2b_path)),
            ("Tax tolerance (per head, rupees)", str(self.tax_tolerance)),
            ("Invoice-number fuzzy threshold", str(self.invoice_fuzzy_threshold)),
            ("GSTIN fuzzy threshold", str(self.gstin_fuzzy_threshold)),
            ("Books documents ingested", str(self.books_document_count)),
            ("GSTR-2B documents ingested", str(self.gstr2b_document_count)),
        ]


def _style_header(ws: Worksheet, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_bridge_sheet(ws: Worksheet, bridge: ItcBridge) -> None:
    ws.title = "ITC Bridge"
    ws.append(["Line", "Reason Code(s)", "Documents", "Amount (Rs.)"])
    _style_header(ws, 1, 4)

    ws.append(["Total ITC as per Books", "", "", float(bridge.books_total)])
    for line in bridge.lines:
        ws.append([line.label, line.reason_codes, line.document_count, float(line.amount)])
    ws.append(
        ["Total ITC as per GSTR-2B (available)", "", "", float(bridge.gstr2b_total_available)]
    )
    ws.append(
        [
            "ITC blocked per GSTR-2B (informational, Section 17(5) etc.)",
            "",
            "",
            float(bridge.gstr2b_total_blocked),
        ]
    )

    closing_row = ws.max_row + 1
    ws.cell(row=closing_row, column=1, value="Closing variance (must be nil)").font = _TOTAL_FONT
    ws.cell(row=closing_row, column=4, value=float(bridge.closing_variance)).font = _TOTAL_FONT

    _autosize(ws, [55, 30, 12, 18])


_SEVERITY_FILL = {
    "high": PatternFill(start_color="FCA5A5", end_color="FCA5A5", fill_type="solid"),
    "medium": PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid"),
    "low": PatternFill(start_color="BBF7D0", end_color="BBF7D0", fill_type="solid"),
}


def _write_register_sheet(ws: Worksheet, findings: list[Exception_]) -> None:
    ws.title = "Exception Register"
    headers = [
        "Reason Code",
        "Severity",
        "Description",
        "Match Category",
        "Books Row Ref",
        "Books Invoice No",
        "GSTR-2B Row Ref",
        "GSTR-2B Invoice No",
        "Supplier GSTIN",
        "ITC at Risk (Rs.)",
        "Evidence",
    ]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    for finding in findings:
        books, gstr2b = finding.books_doc, finding.gstr2b_doc
        anchor = books if books is not None else gstr2b
        gstin = anchor.supplier_gstin if anchor is not None else ""
        ws.append(
            [
                finding.reason_code.value,
                finding.severity.value,
                finding.description,
                finding.match_category.value,
                books.row_ref if books else "",
                books.invoice_number if books else "",
                gstr2b.row_ref if gstr2b else "",
                gstr2b.invoice_number if gstr2b else "",
                gstin,
                float(finding.itc_at_risk),
                finding.evidence,
            ]
        )
        ws.cell(row=ws.max_row, column=2).fill = _SEVERITY_FILL[finding.severity.value]

    _autosize(ws, [10, 10, 45, 18, 22, 18, 22, 18, 18, 15, 60])


def _write_reproducibility_sheet(ws: Worksheet, metadata: RunMetadata) -> None:
    ws.title = "Reproducibility"
    ws.append(["Field", "Value"])
    _style_header(ws, 1, 2)
    for label, value in metadata.to_rows():
        ws.append([label, value])
    _autosize(ws, [35, 70])


def write_working_paper(
    output_path: str | Path,
    *,
    books_docs: list[Document],
    gstr2b_docs: list[Document],
    reconciliation: ReconciliationResult,
    metadata: RunMetadata,
) -> ItcBridge:
    """Write the three-sheet working paper and return the computed bridge
    (mainly so the CLI can print the closing variance without re-reading the
    file it just wrote)."""
    bridge = build_itc_bridge(books_docs, gstr2b_docs, reconciliation)

    wb = Workbook()
    bridge_sheet = wb.active
    assert bridge_sheet is not None
    _write_bridge_sheet(bridge_sheet, bridge)
    _write_register_sheet(wb.create_sheet(), reconciliation.findings)
    _write_reproducibility_sheet(wb.create_sheet(), metadata)

    wb.save(str(output_path))
    return bridge
