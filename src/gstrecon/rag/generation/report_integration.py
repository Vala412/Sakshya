"""Append an "AI Explanations" sheet to a working paper.

Deliberately not a change to report.py: the core engine (ingest -> match ->
classify -> report) has zero dependency on OpenAI/Qdrant, and importing this
module from report.py would break that. This module imports report.py's
public write_working_paper() output (a plain openpyxl Workbook), not the
other way around -- the dependency only ever points from the RAG layer
toward the core engine, never back.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from gstrecon.exceptions import Exception_
from gstrecon.rag.generation.explain import ExplanationResult

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_UNGROUNDED_FILL = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")


def write_explanations_sheet(
    ws: Worksheet,
    findings: list[Exception_],
    explanations: list[ExplanationResult | None],
) -> None:
    """`findings` and `explanations` must be the same list and its
    `explain_findings()` output, in the same order -- this only formats,
    it doesn't re-derive or re-order anything, so a finding whose reason
    code was out of scope (explanation is None) is skipped rather than
    guessed at.
    """
    ws.title = "AI Explanations"
    headers = ["Reason Code", "Grounded", "Explanation", "Citations"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for finding, explanation in zip(findings, explanations, strict=True):
        if explanation is None:
            continue
        citations = "; ".join(
            f"{c.citation} (score {c.score:.2f})" for c in explanation.citations
        )
        ws.append(
            [
                finding.reason_code.value,
                "Yes" if explanation.grounded else "No",
                explanation.explanation,
                citations,
            ]
        )
        if not explanation.grounded:
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = _UNGROUNDED_FILL

    widths = [10, 10, 70, 60]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
        ws.cell(row=1, column=idx).alignment = Alignment(wrap_text=False)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
