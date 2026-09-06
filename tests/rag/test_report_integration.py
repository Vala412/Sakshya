from decimal import Decimal

from openpyxl import Workbook

from gstrecon.exceptions import Exception_, ReasonCode, Severity
from gstrecon.matching import MatchCategory
from gstrecon.rag.generation.explain import Citation, ExplanationResult
from gstrecon.rag.generation.report_integration import write_explanations_sheet


def _finding(reason_code: ReasonCode) -> Exception_:
    return Exception_(
        reason_code=reason_code,
        severity=Severity.HIGH,
        description="test description",
        match_category=MatchCategory.UNMATCHED,
        books_doc=None,
        gstr2b_doc=None,
        itc_at_risk=Decimal("100.00"),
        evidence="test evidence",
    )


def _grounded(reason_code: ReasonCode) -> ExplanationResult:
    return ExplanationResult(
        reason_code=reason_code,
        explanation="A grounded explanation.",
        citations=[Citation(citation="CGST Act, Section 17", score=0.63)],
        grounded=True,
    )


def _ungrounded(reason_code: ReasonCode) -> ExplanationResult:
    return ExplanationResult(
        reason_code=reason_code,
        explanation="No supporting reference located.",
        citations=[],
        grounded=False,
    )


class TestWriteExplanationsSheet:
    def test_writes_header_and_grounded_row(self) -> None:
        wb = Workbook()
        ws = wb.active
        findings = [_finding(ReasonCode.EX_09)]
        explanations = [_grounded(ReasonCode.EX_09)]

        write_explanations_sheet(ws, findings, explanations)

        assert ws.title == "AI Explanations"
        assert [c.value for c in ws[1]] == ["Reason Code", "Grounded", "Explanation", "Citations"]
        row = [c.value for c in ws[2]]
        assert row[0] == "EX-09"
        assert row[1] == "Yes"
        assert row[2] == "A grounded explanation."
        assert "Section 17" in row[3]
        assert "0.63" in row[3]

    def test_ungrounded_row_says_no(self) -> None:
        wb = Workbook()
        ws = wb.active
        findings = [_finding(ReasonCode.EX_09)]
        explanations = [_ungrounded(ReasonCode.EX_09)]

        write_explanations_sheet(ws, findings, explanations)

        row = [c.value for c in ws[2]]
        assert row[1] == "No"
        assert row[3] == ""

    def test_out_of_scope_finding_is_skipped_not_blank_row(self) -> None:
        wb = Workbook()
        ws = wb.active
        findings = [_finding(ReasonCode.EX_02), _finding(ReasonCode.EX_09)]
        explanations = [None, _grounded(ReasonCode.EX_09)]

        write_explanations_sheet(ws, findings, explanations)

        # Header + exactly one data row (EX-02's None is skipped, not
        # written as an empty row that would confuse a reviewer scanning
        # the sheet).
        assert ws.max_row == 2
        assert ws.cell(row=2, column=1).value == "EX-09"

    def test_mismatched_list_lengths_raise(self) -> None:
        wb = Workbook()
        ws = wb.active
        import pytest

        with pytest.raises(ValueError):
            write_explanations_sheet(ws, [_finding(ReasonCode.EX_09)], [])
