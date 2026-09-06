import random
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from gstrecon import exceptions, report
from gstrecon.models import DocSource, DocType, Document, ItcAvailability, TaxAmounts

GSTIN = "27AAPFU0939F1ZV"


def _doc(
    source: DocSource,
    doc_type: DocType = DocType.INVOICE,
    invoice_number: str = "INV/0001",
    invoice_date: date = date(2024, 4, 1),
    taxable_value: Decimal = Decimal("1000.00"),
    igst: Decimal = Decimal("180.00"),
    itc_availability: ItcAvailability | None = None,
    row_ref: str = "row",
) -> Document:
    return Document(
        source=source,
        doc_type=doc_type,
        supplier_gstin=GSTIN,
        invoice_number=invoice_number,
        invoice_date=invoice_date,
        taxable_value=taxable_value,
        tax=TaxAmounts(igst=igst),
        itc_availability=itc_availability,
        row_ref=row_ref,
    )


def _books(**kwargs: object) -> Document:
    return _doc(DocSource.BOOKS, **kwargs)  # type: ignore[arg-type]


def _gstr2b(**kwargs: object) -> Document:
    return _doc(DocSource.GSTR_2B, **kwargs)  # type: ignore[arg-type]


class TestItcBridgeTiesToNil:
    def test_clean_exact_match(self) -> None:
        books, gstr2b = [_books()], [_gstr2b()]
        recon = exceptions.reconcile_full(books, gstr2b)
        bridge = report.build_itc_bridge(books, gstr2b, recon)
        assert bridge.closing_variance == Decimal("0.00")

    def test_orphan_on_each_side(self) -> None:
        books = [_books(row_ref="b1"), _books(invoice_number="INV/0002", row_ref="b2")]
        gstr2b = [_gstr2b(invoice_number="INV/9999", row_ref="g1")]
        recon = exceptions.reconcile_full(books, gstr2b)
        bridge = report.build_itc_bridge(books, gstr2b, recon)
        assert bridge.closing_variance == Decimal("0.00")

    def test_credit_notes_and_blocked_itc(self) -> None:
        books = [
            _books(row_ref="b1"),
            _books(doc_type=DocType.CREDIT_NOTE, invoice_number="CN/0001", row_ref="b2"),
        ]
        gstr2b = [
            _gstr2b(row_ref="g1", itc_availability=ItcAvailability.NOT_AVAILABLE),
            _gstr2b(invoice_number="CN/9999", doc_type=DocType.CREDIT_NOTE, row_ref="g2"),
        ]
        recon = exceptions.reconcile_full(books, gstr2b)
        bridge = report.build_itc_bridge(books, gstr2b, recon)
        assert bridge.closing_variance == Decimal("0.00")

    def test_duplicate_bookings_on_both_sides(self) -> None:
        books = [_books(row_ref="b1"), _books(row_ref="b2")]
        gstr2b = [_gstr2b(row_ref="g1"), _gstr2b(row_ref="g2")]
        recon = exceptions.reconcile_full(books, gstr2b)
        bridge = report.build_itc_bridge(books, gstr2b, recon)
        assert bridge.closing_variance == Decimal("0.00")

    def test_value_and_tax_mismatches(self) -> None:
        books = [_books(taxable_value=Decimal("1000.00"), igst=Decimal("180.00"))]
        gstr2b = [_gstr2b(taxable_value=Decimal("1200.00"), igst=Decimal("216.00"))]
        recon = exceptions.reconcile_full(books, gstr2b)
        bridge = report.build_itc_bridge(books, gstr2b, recon)
        assert bridge.closing_variance == Decimal("0.00")

    def test_bridge_ties_to_nil_property(self) -> None:
        """Randomized check: across many random document sets, the bridge's
        closing variance is always exactly zero. This is what makes the
        bridge a decomposition rather than a hand-signed guess -- a single
        counterexample here would mean the arithmetic identity doesn't hold,
        which is a correctness bug, not an audit finding.
        """
        rng = random.Random(42)
        invoice_pool = [f"INV/{i:04d}" for i in range(1, 9)]

        for trial in range(50):
            books_docs = []
            gstr2b_docs = []
            for i in range(rng.randint(1, 6)):
                doc_type = rng.choice([DocType.INVOICE, DocType.CREDIT_NOTE])
                inv_no = rng.choice(invoice_pool)
                taxable = Decimal(rng.randint(100, 5000))
                igst = (taxable * Decimal("18") / Decimal("100")).quantize(Decimal("0.01"))
                itc_avail = rng.choice(
                    [ItcAvailability.AVAILABLE, ItcAvailability.NOT_AVAILABLE, None]
                )
                books_docs.append(
                    _books(
                        doc_type=doc_type,
                        invoice_number=inv_no,
                        taxable_value=taxable,
                        igst=igst,
                        row_ref=f"t{trial}-b{i}",
                    )
                )
                if rng.random() < 0.7:
                    # Perturb the 2B side sometimes so pairs, orphans, and
                    # value mismatches all show up across trials.
                    taxable_2b = (
                        taxable if rng.random() < 0.5 else taxable + Decimal(rng.randint(-50, 50))
                    )
                    igst_2b = igst if rng.random() < 0.5 else igst + Decimal(rng.randint(-10, 10))
                    gstr2b_docs.append(
                        _gstr2b(
                            doc_type=doc_type,
                            invoice_number=inv_no,
                            taxable_value=taxable_2b,
                            igst=igst_2b,
                            itc_availability=itc_avail,
                            row_ref=f"t{trial}-g{i}",
                        )
                    )

            recon = exceptions.reconcile_full(books_docs, gstr2b_docs)
            bridge = report.build_itc_bridge(books_docs, gstr2b_docs, recon)
            assert bridge.closing_variance == Decimal("0.00"), (
                f"trial {trial} failed to tie out: {bridge.closing_variance}"
            )


class TestWriteWorkingPaper:
    def test_writes_three_sheets_with_expected_content(self, tmp_path: Path) -> None:
        books = [_books(row_ref="b1")]
        gstr2b = [_gstr2b(row_ref="g1", taxable_value=Decimal("1200.00"))]
        recon = exceptions.reconcile_full(books, gstr2b)

        books_file = tmp_path / "books.csv"
        gstr2b_file = tmp_path / "gstr2b.json"
        books_file.write_text("dummy")
        gstr2b_file.write_text("dummy")

        metadata = report.RunMetadata(
            books_path=books_file,
            gstr2b_path=gstr2b_file,
            tax_tolerance=Decimal("0"),
            invoice_fuzzy_threshold=88.0,
            gstin_fuzzy_threshold=85.0,
            books_document_count=len(books),
            gstr2b_document_count=len(gstr2b),
        )

        output_path = tmp_path / "working_paper.xlsx"
        bridge = report.write_working_paper(
            output_path,
            books_docs=books,
            gstr2b_docs=gstr2b,
            reconciliation=recon,
            metadata=metadata,
        )

        assert output_path.exists()
        assert bridge.closing_variance == Decimal("0.00")

        wb = load_workbook(output_path)
        assert wb.sheetnames == ["ITC Bridge", "Exception Register", "Reproducibility"]

        register_ws = wb["Exception Register"]
        assert register_ws.cell(row=1, column=1).value == "Reason Code"
        assert register_ws.cell(row=2, column=1).value == "EX-03"

        repro_ws = wb["Reproducibility"]
        repro_values = {
            repro_ws.cell(row=r, column=1).value for r in range(2, repro_ws.max_row + 1)
        }
        assert "gstrecon version" in repro_values
        assert "Books file SHA-256" in repro_values
